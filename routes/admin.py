"""Admin routes voor KSA Bar."""
import os, io, json, time, uuid
from datetime import datetime
from functools import wraps
from flask import (Blueprint, render_template, request, redirect,
                   url_for, session, send_file, jsonify, flash)
from database.db import get_db, query, execute, add_log, get_setting, set_setting
from config import UPLOADS_PERSONS_DIR, UPLOADS_PRODUCTS_DIR, DATABASE_PATH

admin_bp = Blueprint('admin', __name__)


# ─── Auth ────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def wrap(*a, **kw):
        if not session.get('admin_logged_in'):
            return redirect(url_for('admin.login'))
        logout_min = int(get_setting('admin_logout_min', '10'))
        if time.time() - session.get('admin_last_active', 0) > logout_min * 60:
            session.pop('admin_logged_in', None)
            return redirect(url_for('kiosk.home'))
        session['admin_last_active'] = time.time()
        return f(*a, **kw)
    return wrap


@admin_bp.route('/login', methods=['GET', 'POST'])
def login():
    # Als al ingelogd en sessie nog geldig → direct naar dashboard
    if session.get('admin_logged_in'):
        logout_min = int(get_setting('admin_logout_min', '10'))
        if time.time() - session.get('admin_last_active', 0) <= logout_min * 60:
            session['admin_last_active'] = time.time()
            return redirect(url_for('admin.dashboard'))
    if request.method == 'POST':
        if request.form.get('wachtwoord') == get_setting('admin_password', 'admin123'):
            session['admin_logged_in'] = True
            session['admin_last_active'] = time.time()
            add_log('admin', 'Admin ingelogd')
            return redirect(url_for('admin.dashboard'))
        return render_template('admin/login.html', fout='Incorrect wachtwoord')
    return render_template('admin/login.html')


@admin_bp.route('/logout')
def logout():
    session.pop('admin_logged_in', None)
    add_log('admin', 'Admin uitgelogd')
    return redirect(url_for('kiosk.home'))


# ─── Dashboard ───────────────────────────────────────────────────────────────

@admin_bp.route('/')
@login_required
def dashboard():
    openstaand = query(
        """SELECT p.id, p.voornaam, p.bijnaam,
               COALESCE(SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot),0) as totaal
           FROM persons p
           JOIN order_items oi ON oi.person_id = p.id
           JOIN orders o ON oi.order_id = o.id
           WHERE o.geannuleerd = 0
           GROUP BY p.id HAVING totaal > 0 ORDER BY totaal DESC"""
    )
    stock = query(
        """SELECT p.naam, p.stock, c.naam as categorie
           FROM products p LEFT JOIN categories c ON p.categorie_id=c.id
           WHERE p.actief=1 ORDER BY p.stock ASC LIMIT 20"""
    )
    recente = query(
        """SELECT o.id, o.tijdstip, o.type, p.voornaam, p.bijnaam,
               COUNT(oi.id) as items,
               SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as totaal
           FROM orders o
           LEFT JOIN persons p ON o.gestart_door_id=p.id
           LEFT JOIN order_items oi ON oi.order_id=o.id
           WHERE o.geannuleerd=0
           GROUP BY o.id ORDER BY o.tijdstip DESC LIMIT 10"""
    )
    return render_template('admin/dashboard.html',
                           schulden=openstaand, stock=stock, recente=recente)


# ─── Producten ───────────────────────────────────────────────────────────────

@admin_bp.route('/producten')
@login_required
def producten():
    cat_filter = request.args.get('cat', type=int)
    if cat_filter:
        prods = query(
            """SELECT p.*, c.naam as cat_naam, GROUP_CONCAT(pd.deur) as deuren
               FROM products p
               LEFT JOIN categories c ON p.categorie_id=c.id
               LEFT JOIN product_doors pd ON p.id=pd.product_id
               WHERE p.categorie_id=?
               GROUP BY p.id ORDER BY c.volgorde, p.naam""",
            (cat_filter,)
        )
    else:
        prods = query(
            """SELECT p.*, c.naam as cat_naam, GROUP_CONCAT(pd.deur) as deuren
               FROM products p
               LEFT JOIN categories c ON p.categorie_id=c.id
               LEFT JOIN product_doors pd ON p.id=pd.product_id
               GROUP BY p.id ORDER BY c.volgorde, p.naam"""
        )
    cats = query("SELECT * FROM categories ORDER BY volgorde")
    return render_template('admin/products.html', producten=prods, categorieen=cats)


@admin_bp.route('/producten/nieuw', methods=['POST'])
@login_required
def product_nieuw():
    naam     = request.form.get('naam', '').strip()
    cat_id   = request.form.get('categorie_id', type=int)
    prijs    = request.form.get('verkoop_prijs', 0.0, type=float)
    deuren   = request.form.getlist('deuren')
    bak_grootte_str = request.form.get('bak_grootte', '').strip()
    bak_grootte = int(bak_grootte_str) if bak_grootte_str.isdigit() and int(bak_grootte_str) > 0 else None

    foto_path = _save_foto(request.files.get('foto'), UPLOADS_PRODUCTS_DIR, 'products')

    pid = execute(
        "INSERT INTO products (naam, categorie_id, verkoop_prijs, aankoop_prijs, bak_grootte, foto_path) VALUES (?,?,?,0,?,?)",
        (naam, cat_id, prijs, bak_grootte, foto_path)
    )
    for d in deuren:
        try:
            execute("INSERT INTO product_doors (product_id,deur) VALUES (?,?)", (pid, int(d)))
        except Exception:
            pass
    add_log('admin', f"Product aangemaakt: {naam}")
    return redirect(url_for('admin.producten'))


@admin_bp.route('/producten/<int:pid>/bewerken', methods=['POST'])
@login_required
def product_bewerken(pid):
    old = query("SELECT naam, foto_path, stock, aankoop_prijs, verkoop_prijs FROM products WHERE id=?", (pid,), one=True)
    foto_path = _save_foto(request.files.get('foto'), UPLOADS_PRODUCTS_DIR, 'products') \
                or (old['foto_path'] if old else '')
    naam  = request.form.get('naam', '').strip()
    cat   = request.form.get('categorie_id', type=int)
    prijs = request.form.get('verkoop_prijs', type=float)
    actief= 1 if request.form.get('actief') else 0
    deuren= request.form.getlist('deuren')
    bak_grootte_str = request.form.get('bak_grootte', '').strip()
    bak_grootte = int(bak_grootte_str) if bak_grootte_str.isdigit() and int(bak_grootte_str) > 0 else None
    # Aankoopprijs en stock worden NIET aangepast via admin productenbeheer
    huidige_stock    = old['stock']    if old else 0
    huidige_aankoop  = old['aankoop_prijs'] if old else 0

    execute("UPDATE products SET naam=?,categorie_id=?,verkoop_prijs=?,aankoop_prijs=?,actief=?,stock=?,foto_path=?,bak_grootte=? WHERE id=?",
            (naam, cat, prijs, huidige_aankoop, actief, huidige_stock, foto_path, bak_grootte, pid))
    execute("DELETE FROM product_doors WHERE product_id=?", (pid,))
    for d in deuren:
        try:
            execute("INSERT INTO product_doors (product_id,deur) VALUES (?,?)", (pid, int(d)))
        except Exception:
            pass
    wijzigingen = []
    if old and old['naam'] != naam:
        wijzigingen.append(f"naam: {old['naam']} → {naam}")
    if old and prijs is not None and round(float(old['verkoop_prijs'] or 0), 2) != round(prijs, 2):
        wijzigingen.append(f"prijs: €{float(old['verkoop_prijs'] or 0):.2f} → €{prijs:.2f}")
    detail = f" ({', '.join(wijzigingen)})" if wijzigingen else ""
    add_log('admin', f"Product #{pid} bijgewerkt: {naam}{detail}")
    return redirect(url_for('admin.producten'))


@admin_bp.route('/producten/<int:pid>/verwijderen', methods=['POST'])
@login_required
def product_verwijderen(pid):
    p = query("SELECT naam FROM products WHERE id=?", (pid,), one=True)
    execute("UPDATE products SET actief=0 WHERE id=?", (pid,))
    add_log('admin', f"Product verwijderd (soft): {p['naam'] if p else pid}")
    return redirect(url_for('admin.producten'))


# ─── Categorieën ─────────────────────────────────────────────────────────────

@admin_bp.route('/categorieen/nieuw', methods=['POST'])
@login_required
def categorie_nieuw():
    naam = request.form.get('naam', '').strip()
    if naam:
        execute("INSERT OR IGNORE INTO categories (naam, volgorde) "
                "VALUES (?, (SELECT COALESCE(MAX(volgorde),0)+1 FROM categories))", (naam,))
        add_log('admin', f"Categorie aangemaakt: {naam}")
    return redirect(url_for('admin.producten'))


@admin_bp.route('/categorieen/<int:cid>/bewerken', methods=['POST'])
@login_required
def categorie_bewerken(cid):
    naam    = request.form.get('naam', '').strip()
    volgorde = request.form.get('volgorde', type=int)
    if naam:
        execute("UPDATE categories SET naam=?, volgorde=? WHERE id=?", (naam, volgorde, cid))
        add_log('admin', f"Categorie #{cid} bijgewerkt: {naam}")
    return redirect(url_for('admin.producten'))


@admin_bp.route('/personen/<int:pid>/toggle', methods=['POST'])
@login_required
def persoon_toggle(pid):
    p = query("SELECT actief, is_bond FROM persons WHERE id=?", (pid,), one=True)
    if p and not p['is_bond']:
        nieuw = 0 if p['actief'] else 1
        execute("UPDATE persons SET actief=? WHERE id=?", (nieuw, pid))
        add_log('admin', f"Persoon #{pid} {'geactiveerd' if nieuw else 'gedeactiveerd'}")
    return redirect(url_for('admin.personen'))


@admin_bp.route('/categorieen/<int:cid>/verwijderen', methods=['POST'])
@login_required
def categorie_verwijderen(cid):
    c = query("SELECT naam FROM categories WHERE id=?", (cid,), one=True)
    # Ontkoppel producten van deze categorie voor verwijdering
    execute("UPDATE products SET categorie_id=NULL WHERE categorie_id=?", (cid,))
    execute("DELETE FROM person_blocked_categories WHERE category_id=?", (cid,))
    execute("DELETE FROM categories WHERE id=?", (cid,))
    add_log('admin', f"Categorie verwijderd: {c['naam'] if c else cid}")
    return redirect(url_for('admin.producten'))


# ─── Personen ─────────────────────────────────────────────────────────────────

@admin_bp.route('/personen')
@login_required
def personen():
    p = query("SELECT * FROM persons ORDER BY is_bond DESC, voornaam")
    return render_template('admin/persons.html', personen=p)


@admin_bp.route('/personen/nieuw', methods=['POST'])
@login_required
def persoon_nieuw():
    voornaam  = request.form.get('voornaam', '').strip()
    achternaam = request.form.get('achternaam', '').strip()
    bijnaam   = request.form.get('bijnaam', '').strip()
    if not voornaam:
        return redirect(url_for('admin.dashboard'))
    foto_path = _save_foto(request.files.get('foto'), UPLOADS_PERSONS_DIR, 'persons')
    execute(
        "INSERT INTO persons (voornaam, achternaam, bijnaam, foto_path) VALUES (?,?,?,?)",
        (voornaam, achternaam, bijnaam, foto_path or '')
    )
    add_log('admin', f"Persoon aangemaakt: {bijnaam or voornaam}")
    return redirect(url_for('admin.personen'))


@admin_bp.route('/personen/<int:pid>/bewerken', methods=['POST'])
@login_required
def persoon_bewerken(pid):
    p = query("SELECT * FROM persons WHERE id=?", (pid,), one=True)
    if not p or p['is_bond']:
        return redirect(url_for('admin.personen'))

    foto_path = _save_foto(request.files.get('foto'), UPLOADS_PERSONS_DIR, 'persons') \
                or p['foto_path']
    nieuw_voornaam   = request.form.get('voornaam', '').strip()
    nieuw_achternaam = request.form.get('achternaam', '').strip()
    nieuw_bijnaam    = request.form.get('bijnaam', '').strip()
    nieuw_actief     = 1 if request.form.get('actief') else 0
    execute("UPDATE persons SET voornaam=?,achternaam=?,bijnaam=?,actief=?,foto_path=? WHERE id=?",
            (nieuw_voornaam, nieuw_achternaam, nieuw_bijnaam, nieuw_actief, foto_path, pid))
    wijzigingen = []
    if p['voornaam'] != nieuw_voornaam:
        wijzigingen.append(f"voornaam: {p['voornaam']} → {nieuw_voornaam}")
    if (p['bijnaam'] or '') != nieuw_bijnaam:
        wijzigingen.append(f"bijnaam: {p['bijnaam'] or ''} → {nieuw_bijnaam}")
    if p['actief'] != nieuw_actief:
        wijzigingen.append(f"actief: {'ja' if p['actief'] else 'nee'} → {'ja' if nieuw_actief else 'nee'}")
    detail = f" ({', '.join(wijzigingen)})" if wijzigingen else ""
    add_log('admin', f"Persoon #{pid} bijgewerkt{detail}", pid)
    return redirect(url_for('admin.personen'))


@admin_bp.route('/personen/<int:pid>/verwijderen', methods=['POST'])
@login_required
def persoon_verwijderen(pid):
    p = query("SELECT * FROM persons WHERE id=?", (pid,), one=True)
    if not p or p['is_bond']:
        return redirect(url_for('admin.personen'))
    execute("UPDATE persons SET actief=0 WHERE id=?", (pid,))
    add_log('admin', f"Persoon #{pid} gedeactiveerd")
    return redirect(url_for('admin.personen'))


@admin_bp.route('/personen/<int:pid>/beperkingen', methods=['GET', 'POST'])
@login_required
def persoon_beperkingen(pid):
    p = query("SELECT * FROM persons WHERE id=?", (pid,), one=True)
    if not p or p['is_bond']:
        return redirect(url_for('admin.personen'))

    if request.method == 'POST':
        # Geblokkeerde categorieën
        execute("DELETE FROM person_blocked_categories WHERE person_id=?", (pid,))
        for cid in request.form.getlist('blocked_cats'):
            try:
                execute("INSERT INTO person_blocked_categories (person_id,category_id) VALUES (?,?)",
                        (pid, int(cid)))
            except Exception:
                pass
        # Geblokkeerde producten
        execute("DELETE FROM person_blocked_products WHERE person_id=?", (pid,))
        for prod_id in request.form.getlist('blocked_prods'):
            try:
                execute("INSERT INTO person_blocked_products (person_id,product_id) VALUES (?,?)",
                        (pid, int(prod_id)))
            except Exception:
                pass
        add_log('admin', f"Beperkingen bijgewerkt voor persoon #{pid}")
        return redirect(url_for('admin.persoon_beperkingen', pid=pid))

    cats  = query("SELECT * FROM categories ORDER BY volgorde")
    prods = query(
        """SELECT p.*, c.naam as cat_naam FROM products p
           LEFT JOIN categories c ON p.categorie_id=c.id
           WHERE p.actief=1 ORDER BY c.volgorde, p.naam"""
    )
    blocked_cats  = {r['category_id'] for r in
                     query("SELECT category_id FROM person_blocked_categories WHERE person_id=?", (pid,))}
    blocked_prods = {r['product_id'] for r in
                     query("SELECT product_id FROM person_blocked_products WHERE person_id=?", (pid,))}
    return render_template('admin/person_restrictions.html',
                           persoon=p, categorieen=cats, producten=prods,
                           blocked_cats=blocked_cats, blocked_prods=blocked_prods)


# ─── Bestellingen ─────────────────────────────────────────────────────────────

@admin_bp.route('/bestellingen')
@login_required
def bestellingen():
    filters = {
        'datum_van':  request.args.get('datum_van', ''),
        'datum_tot':  request.args.get('datum_tot', ''),
        'person_id':  request.args.get('person_id', type=int),
    }
    sql = """SELECT o.id, o.tijdstip, o.type, o.geannuleerd, o.video_path,
                    p.voornaam, p.bijnaam,
                    COUNT(oi.id) as items,
                    SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as totaal
             FROM orders o
             LEFT JOIN persons p ON o.gestart_door_id=p.id
             LEFT JOIN order_items oi ON oi.order_id=o.id
             WHERE 1=1"""
    params = []
    if filters['datum_van']:
        sql += " AND o.tijdstip >= ?"; params.append(filters['datum_van'])
    if filters['datum_tot']:
        sql += " AND o.tijdstip <= ?"; params.append(filters['datum_tot'] + ' 23:59:59')
    if filters['person_id']:
        sql += " AND oi.person_id = ?"; params.append(filters['person_id'])
    sql += " GROUP BY o.id ORDER BY o.tijdstip DESC LIMIT 200"

    pers = query("SELECT id,voornaam,bijnaam FROM persons WHERE actief=1 ORDER BY voornaam")
    return render_template('admin/orders.html',
                           bestellingen=query(sql, params),
                           personen=pers, filters=filters)


@admin_bp.route('/bestellingen/<int:oid>')
@login_required
def bestelling_detail(oid):
    order = query("SELECT * FROM orders WHERE id=?", (oid,), one=True)
    items = query(
        """SELECT oi.*, p.voornaam, p.bijnaam, pr.naam as product_naam
           FROM order_items oi
           JOIN persons p  ON oi.person_id=p.id
           JOIN products pr ON oi.product_id=pr.id
           WHERE oi.order_id=?""", (oid,)
    )
    return render_template('admin/order_detail.html', order=order, items=items)


@admin_bp.route('/bestellingen/<int:oid>/bewerken', methods=['POST'])
@login_required
def bestelling_bewerken(oid):
    from services.fifo import consume_stock, restore_stock
    items = query("SELECT * FROM order_items WHERE order_id=?", (oid,))
    for item in items:
        nieuwe_qty = request.form.get(f"qty_{item['id']}", type=int)
        if nieuwe_qty is None:
            continue
        verschil = nieuwe_qty - item['hoeveelheid']
        if nieuwe_qty == 0:
            restore_stock(item['product_id'], item['hoeveelheid'])
            execute("DELETE FROM order_items WHERE id=?", (item['id'],))
        elif verschil > 0:
            consume_stock(item['product_id'], verschil)
            execute("UPDATE order_items SET hoeveelheid=? WHERE id=?", (nieuwe_qty, item['id']))
        elif verschil < 0:
            restore_stock(item['product_id'], abs(verschil))
            execute("UPDATE order_items SET hoeveelheid=? WHERE id=?", (nieuwe_qty, item['id']))
    add_log('admin', f"Bestelling #{oid} bewerkt")
    return redirect(url_for('admin.bestelling_detail', oid=oid))


@admin_bp.route('/bestellingen/<int:oid>/verwijderen', methods=['POST'])
@login_required
def bestelling_verwijderen(oid):
    from services.fifo import restore_stock
    items = query("SELECT * FROM order_items WHERE order_id=?", (oid,))
    for item in items:
        restore_stock(item['product_id'], item['hoeveelheid'])
    execute("UPDATE orders SET geannuleerd=1 WHERE id=?", (oid,))
    add_log('admin', f"Bestelling #{oid} geannuleerd")
    return redirect(url_for('admin.bestellingen'))


# ─── Rekening ─────────────────────────────────────────────────────────────────

@admin_bp.route('/rekening')
@login_required
def rekening():
    datum_van = request.args.get('datum_van', datetime.now().replace(day=1).strftime('%Y-%m-%d'))
    datum_tot = request.args.get('datum_tot', datetime.now().strftime('%Y-%m-%d'))

    overzicht = query(
        """SELECT p.id, p.voornaam, p.achternaam, p.bijnaam, p.is_bond,
               COALESCE(SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot),0) as totaal,
               COUNT(DISTINCT o.id) as bestellingen
           FROM persons p
           LEFT JOIN order_items oi ON oi.person_id=p.id
           LEFT JOIN orders o ON oi.order_id=o.id AND o.geannuleerd=0
               AND o.tijdstip BETWEEN ? AND ?
           WHERE p.actief=1
           GROUP BY p.id HAVING totaal > 0
           ORDER BY p.is_bond, totaal DESC""",
        (datum_van, datum_tot + ' 23:59:59')
    )
    return render_template('admin/invoice.html',
                           overzicht=overzicht, datum_van=datum_van, datum_tot=datum_tot)


@admin_bp.route('/rekening/persoon/<int:pid>')
@login_required
def rekening_persoon(pid):
    p = query("SELECT * FROM persons WHERE id=?", (pid,), one=True)
    if not p:
        return redirect(url_for('admin.rekening'))

    # All-time orders totaal
    totaal_orders = query(
        """SELECT COALESCE(SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot), 0) as totaal
           FROM order_items oi
           JOIN orders o ON oi.order_id=o.id
           WHERE oi.person_id=? AND o.geannuleerd=0""",
        (pid,), one=True
    )['totaal']

    # All-time betalingen totaal
    totaal_betalingen = query(
        "SELECT COALESCE(SUM(bedrag), 0) as totaal FROM betalingen WHERE person_id=?",
        (pid,), one=True
    )['totaal']

    saldo = round(totaal_orders - totaal_betalingen, 2)

    # Betalingen geschiedenis
    betalingen = query(
        "SELECT * FROM betalingen WHERE person_id=? ORDER BY tijdstip DESC",
        (pid,)
    )

    # Bestellingen per product (all-time)
    items = query(
        """SELECT pr.naam, SUM(oi.hoeveelheid) as qty,
               oi.verkoop_prijs_snapshot as prijs,
               SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as subtotaal
           FROM order_items oi
           JOIN products pr ON oi.product_id=pr.id
           JOIN orders o ON oi.order_id=o.id
           WHERE oi.person_id=? AND o.geannuleerd=0
           GROUP BY pr.id, oi.verkoop_prijs_snapshot ORDER BY pr.naam""",
        (pid,)
    )

    # Individuele bestellingen
    bestellingen = query(
        """SELECT o.id, o.tijdstip, o.type,
               COUNT(oi.id) as items,
               SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as totaal
           FROM orders o
           JOIN order_items oi ON oi.order_id=o.id
           WHERE oi.person_id=? AND o.geannuleerd=0
           GROUP BY o.id ORDER BY o.tijdstip DESC""",
        (pid,)
    )

    # Winst/verlies: verkoop - aankoop (per besteld product)
    winst_data = query(
        """SELECT SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as omzet,
               SUM(oi.hoeveelheid * pr.aankoop_prijs) as kosten
           FROM order_items oi
           JOIN products pr ON oi.product_id=pr.id
           JOIN orders o ON oi.order_id=o.id
           WHERE oi.person_id=? AND o.geannuleerd=0""",
        (pid,), one=True
    )
    omzet  = winst_data['omzet']  or 0
    kosten = winst_data['kosten'] or 0
    winst  = round(omzet - kosten, 2)

    return render_template('admin/invoice_person.html',
                           persoon=p, items=items,
                           totaal_orders=totaal_orders,
                           totaal_betalingen=totaal_betalingen,
                           saldo=saldo,
                           betalingen=betalingen,
                           bestellingen=bestellingen,
                           omzet=omzet, kosten=kosten, winst=winst)


@admin_bp.route('/rekening/persoon/<int:pid>/betaling', methods=['POST'])
@login_required
def rekening_betaling(pid):
    try:
        bedrag = float(request.form.get('bedrag', 0))
    except ValueError:
        bedrag = 0
    beschrijving = request.form.get('beschrijving', '').strip()
    if bedrag != 0:
        execute(
            "INSERT INTO betalingen (person_id, bedrag, beschrijving) VALUES (?,?,?)",
            (pid, bedrag, beschrijving)
        )
        add_log('betaling', f"Betaling €{bedrag:.2f} — {beschrijving or 'geen opmerking'}",
                pid, None, 'betaling')
        flash(f'Betaling van €{bedrag:.2f} geregistreerd.', 'success')
    return redirect(url_for('admin.rekening_persoon', pid=pid))


@admin_bp.route('/rekening/persoon/<int:pid>/sluiten', methods=['POST'])
@login_required
def rekening_sluiten(pid):
    totaal_orders = query(
        """SELECT COALESCE(SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot), 0) as totaal
           FROM order_items oi JOIN orders o ON oi.order_id=o.id
           WHERE oi.person_id=? AND o.geannuleerd=0""",
        (pid,), one=True
    )['totaal']
    totaal_betalingen = query(
        "SELECT COALESCE(SUM(bedrag), 0) as totaal FROM betalingen WHERE person_id=?",
        (pid,), one=True
    )['totaal']
    saldo = round(totaal_orders - totaal_betalingen, 2)
    if saldo > 0:
        execute(
            "INSERT INTO betalingen (person_id, bedrag, beschrijving) VALUES (?,?,?)",
            (pid, saldo, 'Rekening gesloten')
        )
        add_log('betaling', f"Rekening gesloten — €{saldo:.2f} betaald", pid, None, 'betaling')
        flash(f'Rekening gesloten. €{saldo:.2f} geregistreerd als betaald.', 'success')
    else:
        flash('Saldo is al nul of negatief, niets gesloten.', 'info')
    return redirect(url_for('admin.rekening_persoon', pid=pid))


@admin_bp.route('/rekening/export')
@login_required
def rekening_export():
    datum_van = request.args.get('datum_van', '')
    datum_tot = request.args.get('datum_tot', '')
    formaat   = request.args.get('formaat', 'excel')

    rows = query(
        """SELECT p.voornaam, p.achternaam, p.bijnaam, p.is_bond,
               pr.naam as product_naam,
               SUM(oi.hoeveelheid) as qty,
               oi.verkoop_prijs_snapshot as prijs,
               SUM(oi.hoeveelheid * oi.verkoop_prijs_snapshot) as subtotaal
           FROM order_items oi
           JOIN persons p  ON oi.person_id=p.id
           JOIN products pr ON oi.product_id=pr.id
           JOIN orders o ON oi.order_id=o.id
           WHERE o.geannuleerd=0 AND o.tijdstip BETWEEN ? AND ?
           GROUP BY p.id, pr.id, oi.verkoop_prijs_snapshot
           ORDER BY p.is_bond, p.voornaam, pr.naam""",
        (datum_van, datum_tot + ' 23:59:59')
    )

    def naam(r):
        return r['bijnaam'] or f"{r['voornaam']} {r['achternaam'] or ''}".strip()

    if formaat == 'excel':
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Rekening"
        ws.append(['Naam', 'Product', 'Stuks', 'Prijs/stuk', 'Subtotaal'])
        for r in rows:
            ws.append([naam(r), r['product_naam'], r['qty'],
                       round(r['prijs'], 2), round(r['subtotaal'], 2)])
        out = io.BytesIO()
        wb.save(out); out.seek(0)
        return send_file(out,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"rekening_{datum_van}_{datum_tot}.xlsx")

    elif formaat == 'pdf':
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors

        out = io.BytesIO()
        doc = SimpleDocTemplate(out, pagesize=A4)
        styles = getSampleStyleSheet()
        els = [Paragraph(f"KSA Bar — Rekening {datum_van} t/m {datum_tot}", styles['Title']),
               Spacer(1, 12)]

        data = [['Naam', 'Product', 'Stuks', 'Prijs', 'Subtotaal']]
        for r in rows:
            data.append([naam(r), r['product_naam'], str(r['qty']),
                         f"€{r['prijs']:.2f}", f"€{r['subtotaal']:.2f}"])
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1a1a2e')),
            ('TEXTCOLOR',  (0,0), (-1,0), colors.white),
            ('GRID',       (0,0), (-1,-1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#f0f0f0')]),
        ]))
        els.append(t)
        doc.build(els)
        out.seek(0)
        return send_file(out, mimetype='application/pdf',
                         as_attachment=True,
                         download_name=f"rekening_{datum_van}_{datum_tot}.pdf")



# ─── Baravonden ───────────────────────────────────────────────────────────────

@admin_bp.route('/baravond')
@login_required
def baravond_list():
    baravonden = query(
        """SELECT be.*,
                  pa.voornaam as activator_voornaam, pa.bijnaam as activator_bijnaam,
                  pd.voornaam as deactivator_voornaam, pd.bijnaam as deactivator_bijnaam
           FROM bar_evenings be
           LEFT JOIN persons pa ON be.activator_id = pa.id
           LEFT JOIN persons pd ON be.deactivator_id = pd.id
           ORDER BY be.start_tijd DESC"""
    )
    return render_template('admin/bar_evenings.html', baravonden=baravonden)


@admin_bp.route('/baravond/<int:bid>')
@login_required
def baravond_detail(bid):
    be = query("SELECT * FROM bar_evenings WHERE id=?", (bid,), one=True)
    if not be:
        return redirect(url_for('admin.baravond_list'))

    start_inv = json.loads(be['start_inventaris'] or '{}')
    eind_inv  = json.loads(be['eind_inventaris'] or '{}')
    verbruik  = json.loads(be['verbruik'] or '{}')

    # Verrijken met productnamen + prijzen
    product_ids = set(start_inv) | set(verbruik)
    prods = {}
    if product_ids:
        placeholders = ','.join('?' * len(product_ids))
        rows = query(
            f"SELECT id, naam, verkoop_prijs, aankoop_prijs FROM products WHERE id IN ({placeholders})",
            [int(i) for i in product_ids]
        )
        prods = {str(r['id']): dict(r) for r in rows}

    regels = []
    totaal_omzet = totaal_kosten = 0.0
    for pid_str, verbruikt in verbruik.items():
        if verbruikt <= 0:
            continue
        prod = prods.get(pid_str, {})
        vprijs = prod.get('verkoop_prijs', 0) or 0
        aprijs = prod.get('aankoop_prijs', 0) or 0
        omzet  = verbruikt * vprijs
        kosten = verbruikt * aprijs
        totaal_omzet  += omzet
        totaal_kosten += kosten
        regels.append({
            'naam':      prod.get('naam', f'Product {pid_str}'),
            'start':     start_inv.get(pid_str, 0),
            'eind':      eind_inv.get(pid_str, 0),
            'verbruikt': verbruikt,
            'vprijs':    vprijs,
            'aprijs':    aprijs,
            'omzet':     omzet,
            'kosten':    kosten,
            'winst':     omzet - kosten,
        })
    regels.sort(key=lambda r: r['verbruikt'], reverse=True)

    activator = query("SELECT voornaam, bijnaam FROM persons WHERE id=?",
                      (be['activator_id'],), one=True)
    deactivator = (query("SELECT voornaam, bijnaam FROM persons WHERE id=?",
                          (be['deactivator_id'],), one=True)
                   if be['deactivator_id'] else None)

    return render_template('admin/bar_evening_detail.html',
                           be=be, regels=regels,
                           activator=activator, deactivator=deactivator,
                           totaal_omzet=round(totaal_omzet, 2),
                           totaal_kosten=round(totaal_kosten, 2),
                           totaal_winst=round(totaal_omzet - totaal_kosten, 2))


@admin_bp.route('/baravond/<int:bid>/naam', methods=['POST'])
@login_required
def baravond_update_naam(bid):
    naam = request.form.get('naam', '').strip()
    execute("UPDATE bar_evenings SET naam=? WHERE id=?", (naam, bid))
    flash('Naam opgeslagen.', 'success')
    return redirect(url_for('admin.baravond_detail', bid=bid))


# ─── Winst ────────────────────────────────────────────────────────────────────

@admin_bp.route('/winst')
@login_required
def winst():
    from services.fifo import get_product_profit_stats
    datum_van = request.args.get('datum_van', datetime.now().replace(day=1).strftime('%Y-%m-%d'))
    datum_tot = request.args.get('datum_tot', datetime.now().strftime('%Y-%m-%d'))

    prods = query("SELECT id, naam FROM products ORDER BY naam")
    stats = []
    for p in prods:
        s = get_product_profit_stats(p['id'], datum_van, datum_tot + ' 23:59:59')
        if s['totaal_stuks'] > 0:
            s['naam'] = p['naam']
            stats.append(s)

    return render_template('admin/profit.html', stats=stats,
                           totaal_omzet=sum(s['omzet'] for s in stats),
                           totaal_kosten=sum(s['kosten'] for s in stats),
                           totaal_winst=sum(s['winst'] for s in stats),
                           datum_van=datum_van, datum_tot=datum_tot)


# ─── Logs ─────────────────────────────────────────────────────────────────────

@admin_bp.route('/logs')
@login_required
def logs():
    filters = {
        'type':      request.args.get('type', ''),
        'person_id': request.args.get('person_id', type=int),
        'datum_van': request.args.get('datum_van', ''),
        'datum_tot': request.args.get('datum_tot', ''),
    }
    sql = """SELECT l.*, p.voornaam, p.bijnaam
             FROM logs l LEFT JOIN persons p ON l.person_id=p.id WHERE 1=1"""
    params = []
    if filters['type']:
        sql += " AND l.type=?"; params.append(filters['type'])
    if filters['person_id']:
        sql += " AND l.person_id=?"; params.append(filters['person_id'])
    if filters['datum_van']:
        sql += " AND l.tijdstip>=?"; params.append(filters['datum_van'])
    if filters['datum_tot']:
        sql += " AND l.tijdstip<=?"; params.append(filters['datum_tot'] + ' 23:59:59')
    sql += " ORDER BY l.tijdstip DESC LIMIT 500"

    pers = query("SELECT id,voornaam,bijnaam FROM persons ORDER BY voornaam")
    # Haal video paden op via referentie_id naar orders
    logs_list = query(sql, params)
    # Voeg video_path toe voor bestellingen
    order_videos = {}
    for log in logs_list:
        if log['referentie_id'] and log['referentie_type'] == 'bestelling':
            if log['referentie_id'] not in order_videos:
                o = query("SELECT video_path FROM orders WHERE id=?",
                          (log['referentie_id'],), one=True)
                order_videos[log['referentie_id']] = (o['video_path'] if o else '') or ''
    return render_template('admin/logs.html',
                           logs=logs_list, personen=pers, filters=filters,
                           order_videos=order_videos,
                           log_types=['bestelling','stock','baravond','aanvulling','admin','systeem','deur'])


# ─── Database viewer ──────────────────────────────────────────────────────────

@admin_bp.route('/database')
@login_required
def database_viewer():
    tabellen = ['persons', 'products', 'categories', 'orders', 'order_items',
                'product_doors', 'bar_evenings', 'refill_sessions',
                'shop_purchases', 'shop_purchase_items', 'logs', 'settings', 'fifo_batches']
    tabel = request.args.get('tabel', 'persons')
    if tabel not in tabellen:
        tabel = 'persons'

    try:
        kolommen = query(f"PRAGMA table_info({tabel})")
        kol_namen = [k['name'] for k in kolommen]
        rijen = query(f"SELECT * FROM {tabel} ORDER BY rowid DESC LIMIT 200")
        totaal = query(f"SELECT COUNT(*) as n FROM {tabel}", one=True)['n']
    except Exception as e:
        kol_namen, rijen, totaal = [], [], 0

    counts = {}
    for t in tabellen:
        try:
            counts[t] = query(f"SELECT COUNT(*) as n FROM {t}", one=True)['n']
        except Exception:
            counts[t] = '?'

    return render_template('admin/database.html',
                           tabellen=tabellen, tabel=tabel,
                           kolommen=kol_namen, rijen=rijen,
                           totaal=totaal, counts=counts)


# ─── Hardware test ────────────────────────────────────────────────────────────

@admin_bp.route('/hardware-test')
@login_required
def hardware_test():
    from hardware.gpio_controller import get_fridge_controller
    status = get_fridge_controller().get_status()
    return render_template('admin/hardware_test.html', deur_status=status)


# ─── Instellingen ─────────────────────────────────────────────────────────────

@admin_bp.route('/instellingen', methods=['GET', 'POST'])
@login_required
def instellingen():
    if request.method == 'POST':
        wijzigingen = []
        for k in ['admin_password','deur_timeout_sec','screensaver_timeout_min',
                  'admin_logout_min','video_bewaar_dagen','pi_reboot_tijd','product_kolommen']:
            v = request.form.get(k, '').strip()
            if v:
                oud = get_setting(k, '')
                if oud != v:
                    if k == 'admin_password':
                        wijzigingen.append('wachtwoord gewijzigd')
                    else:
                        wijzigingen.append(f"{k}: {oud} → {v}")
                set_setting(k, v)
        nieuw_prijs = 'true' if request.form.get('prijs_tonen') == '1' else 'false'
        oud_prijs   = get_setting('prijs_tonen', 'false')
        if oud_prijs != nieuw_prijs:
            wijzigingen.append(f"prijs_tonen: {oud_prijs} → {nieuw_prijs}")
        set_setting('prijs_tonen', nieuw_prijs)
        detail = f": {', '.join(wijzigingen)}" if wijzigingen else " (geen wijzigingen)"
        add_log('admin', f"Instellingen gewijzigd{detail}")
        return redirect(url_for('admin.instellingen'))
    all_s = {r['sleutel']: r['waarde'] for r in query("SELECT sleutel,waarde FROM settings")}
    fout = request.args.get('fout')
    succes = request.args.get('succes')
    return render_template('admin/settings.html', settings=all_s, fout=fout, succes=succes)


# ─── Database leegmaken ───────────────────────────────────────────────────────

TRANSACTIONELE_TABELLEN = [
    'order_items', 'orders',
    'bar_evenings',
    'refill_sessions',
    'shop_purchase_items', 'shop_purchases',
    'fifo_batches',
    'door_events',
    'logs',
    'admin_sessions',
]

@admin_bp.route('/instellingen/database-leegmaken', methods=['POST'])
@login_required
def database_leegmaken():
    bevestiging = request.form.get('bevestiging', '').strip()
    if bevestiging != 'LEEGMAKEN':
        return redirect(url_for('admin.instellingen', fout='leegmaken'))

    from services.backup import backup_database
    backup_database()  # automatische backup vóór het legen

    conn = get_db()
    for tabel in TRANSACTIONELE_TABELLEN:
        conn.execute(f"DELETE FROM {tabel}")
    # Reset stock van alle producten naar 0
    conn.execute("UPDATE products SET stock = 0")
    # Reset auto-increment tellers
    conn.execute(
        "DELETE FROM sqlite_sequence WHERE name IN ({})".format(
            ','.join('?' * len(TRANSACTIONELE_TABELLEN))
        ),
        TRANSACTIONELE_TABELLEN
    )
    conn.commit()
    conn.close()

    add_log('admin', 'Database volledig leeggemaakt')
    return redirect(url_for('admin.instellingen', succes='leeggemaakt'))


# ─── Backup download / upload ─────────────────────────────────────────────────

@admin_bp.route('/backup/download')
@login_required
def backup_download():
    from services.backup import backup_database
    pad = backup_database()
    return send_file(pad, as_attachment=True, download_name=os.path.basename(pad))


@admin_bp.route('/backup/upload', methods=['POST'])
@login_required
def backup_upload():
    bestand = request.files.get('backup_bestand')
    if not bestand or not bestand.filename.endswith('.db'):
        flash('Ongeldig bestand. Upload een .db bestand.', 'error')
        return redirect(url_for('admin.instellingen'))

    from services.backup import backup_database
    backup_database()  # maak eerst een backup van huidige data

    import shutil
    shutil.copy2(DATABASE_PATH, DATABASE_PATH + '.voor_upload.bak')
    bestand.save(DATABASE_PATH)
    add_log('admin', f'Database hersteld via upload: {bestand.filename}')
    flash(f'Database hersteld vanuit {bestand.filename}. Herstart de app voor beste resultaten.', 'success')
    return redirect(url_for('admin.instellingen'))


# ─── Helper ───────────────────────────────────────────────────────────────────

def _save_foto(foto_file, upload_dir: str, subfolder: str) -> str:
    """Sla een geüpload fotobestand op. Geeft relatief pad terug of ''."""
    if not foto_file or not foto_file.filename:
        return ''
    ext = foto_file.filename.rsplit('.', 1)[-1].lower()
    if ext not in ('jpg', 'jpeg', 'png', 'webp'):
        return ''
    naam = f"{uuid.uuid4().hex}.{ext}"
    foto_file.save(os.path.join(upload_dir, naam))
    return f"{subfolder}/{naam}"
